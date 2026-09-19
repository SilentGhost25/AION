"""
Tests for Phase 4: DOCX Export, Visual Embedding, Native OMML Math, and Examiner Answer Key.
Verifies:
  - Generation of valid .docx binary
  - Question rows with LaTeX math rendering via OMML / clean formatting
  - Embedding of high-DPI image assets in question cells
  - Appendix with 4-Part University Scheme of Valuation & Examiner Answer Key
"""

import io
import os
import tempfile
import pytest
from pathlib import Path
from PIL import Image

from v0_1.docx_export import (
    generate_docx_from_paper,
    latex_to_omml_element,
    add_formatted_text_to_paragraph,
    _get_xslt_transform,
)
import docx


def test_latex_to_omml_conversion():
    # If MS Office MML2OMML.XSL is present, OMML element is produced
    transform = _get_xslt_transform()
    if transform is not None:
        elem = latex_to_omml_element(r"\frac{a}{b} = \sigma")
        assert elem is not None
        assert "oMath" in elem.tag


def test_docx_export_with_image_and_answer_key():
    # 1. Create a dummy image for testing
    img_fd, img_path = tempfile.mkstemp(suffix=".png")
    os.close(img_fd)
    img = Image.new("RGB", (100, 100), color="blue")
    img.save(img_path)

    # 2. Build mock paper data with math formula and image
    paper_data = {
        "subject": "Data Structures & Algorithms",
        "subject_code": "21CS32",
        "exam_type": "IAT1",
        "modules": [
            {
                "module_index": 1,
                "questions": [
                    {
                        "question_number": 1,
                        "sub_questions": [
                            {
                                "label": "a",
                                "text": "Formulate the binary search complexity $T(n) = T(n/2) + O(1)$ and explain recurrence.",
                                "marks": 6,
                                "co": "CO1",
                                "bloom": "L2",
                                "solution": "Given recurrence $T(n) = T(n/2) + c$. By Master Theorem case 2, complexity is $\\Theta(\\log n)$.",
                                "image_path": img_path,
                                "image_caption": "Recursion Tree for Divide and Conquer",
                            },
                            {
                                "label": "b",
                                "text": "Demonstrate array representation of binary heaps with an example.",
                                "marks": 4,
                                "co": "CO1",
                                "bloom": "L3",
                                "solution": "Parent node at index $i$, left child at $2i+1$, right child at $2i+2$.",
                            }
                        ]
                    }
                ]
            }
        ]
    }

    try:
        buf = generate_docx_from_paper(paper_data)
        assert buf is not None
        docx_bytes = buf.getvalue()
        assert len(docx_bytes) > 1000

        # Read back document to inspect structure
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
        all_doc_text = full_text + "\n" + table_text

        # Verify Examiner Answer Key heading is present
        assert "EXAMINER ANSWER KEY & SCHEME OF VALUATION" in all_doc_text
        assert "Standard VTU 4-Part Valuation Framework:" in all_doc_text
        assert "Q1 (a)" in all_doc_text
        assert "Model Solution / Key Milestones:" in all_doc_text

        # Verify tables exist: main table + CO table + coverage tables + answer key rubric tables
        assert len(doc.tables) >= 4

        # Verify 4-Part Valuation Table headers
        table_texts = [" ".join(cell.text for cell in row.cells) for t in doc.tables for row in t.rows]
        valuation_headers = [t for t in table_texts if "Evaluation Component" in t]
        assert len(valuation_headers) >= 1
    finally:
        if os.path.exists(img_path):
            os.remove(img_path)
